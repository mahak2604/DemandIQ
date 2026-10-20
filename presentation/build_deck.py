"""
Builds the DemandIQ presentation (PowerPoint, .pptx) from the pipeline outputs.

    python presentation/build_deck.py
"""
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
sys.path.insert(0, str(ROOT / "report"))
from build_report import TEAM, load_results  # noqa: E402

FIG = ROOT / "outputs" / "figures"
SHOTS = ROOT / "outputs" / "screenshots"
LOGO = ROOT / "assets" / "nmims_logo.png"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
NAVY_DK = RGBColor(0x16, 0x2B, 0x47)
TEAL = RGBColor(0x2E, 0x8B, 0x8B)
AMBER = RGBColor(0xD9, 0x8E, 0x04)
RED = RGBColor(0xB2, 0x3A, 0x48)
GREEN = RGBColor(0x4C, 0x8C, 0x4A)
GREY = RGBColor(0x8A, 0x8F, 0x98)
TINT = RGBColor(0xEE, 0xF2, 0xF6)
INK = RGBColor(0x1E, 0x24, 0x2C)
MUTED = RGBColor(0x5A, 0x62, 0x70)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
HEAD = "Cambria"
BODY = "Calibri"

W, H = 13.333, 7.5


# ------------------------------------------------------------------ primitives
def text(slide, x, y, w, h, runs, size=16, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, italic=False, spacing=None, name=None):
    """runs: str, or list of paragraphs; each paragraph is str or list of (text, {opts})."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0)
    tf.margin_top = tf.margin_bottom = Inches(0)
    tf.vertical_anchor = anchor
    paras = runs if isinstance(runs, list) else [runs]
    for i, pr in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        segs = pr if isinstance(pr, list) else [(pr, {})]
        for seg, o in segs:
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.italic = o.get("italic", italic)
            f.color.rgb = o.get("color", color)
    return tb


def box(slide, x, y, w, h, fill=TINT, radius=0.08, line=None, shadow=False, name=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line:
        shp.line.color.rgb = line
        shp.line.width = Pt(1)
    else:
        shp.line.fill.background()
    if not shadow:
        shp.shadow.inherit = False
    if name:
        shp.name = name
    shp.text_frame.text = ""
    return shp


def circle(slide, x, y, d, fill, label, size=16, color=WHITE):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = HEAD
    r.font.color.rgb = color
    return shp


def picture(slide, path, x, y, w=None, h=None):
    """Place an image inside the box (x, y, w, h), keeping its aspect ratio, centred."""
    im = Image.open(path)
    ar = im.width / im.height
    if w and h:
        if w / h > ar:
            nw, nh = h * ar, h
        else:
            nw, nh = w, w / ar
        x, y = x + (w - nw) / 2, y + (h - nh) / 2
        w, h = nw, nh
    elif w:
        h = w / ar
    else:
        w = h * ar
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))


def background(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def content_slide(prs, kicker, title, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, WHITE)
    text(s, 0.6, 0.38, 9.5, 0.3, kicker.upper(), size=11, color=TEAL, bold=True, font=BODY, name="Kicker")
    text(s, 0.6, 0.66, 12.1, 0.8, title, size=30, color=NAVY, bold=True, font=HEAD, name="Title")
    picture(s, LOGO, 12.15, 0.3, h=0.62)
    text(s, 0.6, 7.05, 6, 0.25, "DemandIQ  |  NMIMS MPSTME  |  Machine Learning, Sem V", size=9, color=GREY)
    text(s, 12.2, 7.05, 0.55, 0.25, str(n), size=9, color=GREY, align=PP_ALIGN.RIGHT)
    return s


def notes(slide, txt):
    slide.notes_slide.notes_text_frame.text = txt


def style_chart(chart, size=12):
    chart.font.name = BODY
    chart.font.size = Pt(size)
    chart.font.color.rgb = MUTED


# ------------------------------------------------------------------ slides
def build():
    R = load_results()
    dt, dv, st, pol = R["dt"], R["dv"], R["st"], R["pol"]
    gb, rf, ha = dt.loc["Gradient Boosting"], dt.loc["Random Forest"], dt.loc["Historical average"]
    A, B, C = pol.iloc[0], pol.iloc[1], pol.iloc[2]
    gbs = st.loc["Gradient Boosting"]
    ds, inv = R["ds"], R["inv"]

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    n = 0

    # 1 ---------------------------------------------------------------- title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, NAVY)
    box(s, 0.6, 0.55, 1.35, 1.35, fill=WHITE, radius=0.12)
    picture(s, LOGO, 0.68, 0.63, 1.19, 1.19)
    text(s, 2.2, 0.72, 9, 0.4, "SVKM's NMIMS  |  Mukesh Patel School of Technology Management & Engineering",
         size=14, color=RGBColor(0xC9, 0xD6, 0xE3))
    text(s, 2.2, 1.12, 9, 0.4, "Machine Learning  |  B.Tech Semester V  |  AY 2026-27", size=14,
         color=RGBColor(0xC9, 0xD6, 0xE3))
    text(s, 0.6, 2.45, 12, 1.2, "DemandIQ", size=66, color=WHITE, bold=True, font=HEAD)
    text(s, 0.6, 3.6, 11.5, 0.9, "Machine learning for product demand and stockout risk prediction",
         size=26, color=RGBColor(0xF2, 0xC1, 0x6B), font=HEAD)
    text(s, 0.6, 4.35, 11.5, 0.5, "Forecast what will sell  •  warn what will run out  •  say how much to order  •  explain why",
         size=16, color=RGBColor(0xC9, 0xD6, 0xE3))
    for i, (nm, roll) in enumerate(TEAM):
        x = 0.6 + i * 3.05
        box(s, x, 5.45, 2.85, 0.95, fill=NAVY_DK, radius=0.1)
        text(s, x + 0.2, 5.58, 2.5, 0.4, nm, size=16, color=WHITE, bold=True)
        text(s, x + 0.2, 5.95, 2.5, 0.3, roll, size=13, color=RGBColor(0xC9, 0xD6, 0xE3))
    text(s, 0.6, 6.7, 8, 0.3, "Guide: Prof. ____________________", size=13, color=RGBColor(0xC9, 0xD6, 0xE3))
    notes(s, "Good morning. We are presenting DemandIQ, a classical machine-learning system that forecasts demand, "
             "predicts stockouts, recommends order quantities and explains every number. No deep learning is used.")

    # 2 ---------------------------------------------------------------- problem
    n += 1
    s = content_slide(prs, "The problem", "Every store asks one question: how much should we stock?", n + 1)
    for i, (head, col, pts) in enumerate([
        ("Too little stock", RED, ["The shelf runs empty - the sale is lost", "The customer often goes to a competitor",
                                   "Empty shelves damage trust in the store"]),
        ("Too much stock", AMBER, ["Cash is locked up in goods", "Storage and handling costs keep adding up",
                                   "Old stock is discounted or thrown away"])]):
        x = 0.6 + i * 4.15
        box(s, x, 1.75, 3.95, 3.55, fill=TINT)
        circle(s, x + 0.3, 2.0, 0.62, col, "↓" if i == 0 else "↑", size=22)
        text(s, x + 1.1, 2.1, 2.7, 0.5, head, size=21, bold=True, color=INK, font=HEAD)
        text(s, x + 0.3, 2.95, 3.4, 2.2, [[("•  " + p, {})] for p in pts], size=16, color=INK, spacing=9)
    box(s, 9.05, 1.75, 3.7, 3.55, fill=NAVY)
    text(s, 9.35, 2.05, 3.2, 1.2, "8.3%", size=60, bold=True, color=WHITE, font=HEAD)
    text(s, 9.35, 3.25, 3.2, 1.3, "of items are out of stock on an average day in retail worldwide",
         size=16, color=WHITE)
    text(s, 9.35, 4.75, 3.2, 0.3, "Gruen, Corsten & Bharadwaj (2002)", size=11, color=RGBColor(0xC9, 0xD6, 0xE3), italic=True)
    box(s, 0.6, 5.6, 12.15, 1.1, fill=WHITE, line=RGBColor(0xD5, 0xDC, 0xE4))
    text(s, 0.9, 5.78, 11.6, 0.8, [[("Root cause: ", {"bold": True, "color": NAVY}),
                                    ("nobody knows tomorrow's demand. Rules like \"order last month's sales plus a bit\" "
                                     "ignore promotions, holidays and weekdays - so stores run short exactly when demand jumps.", {})]],
         size=16, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "Both mistakes cost money. A worldwide study found 8.3% of items out of stock on a normal day. The root "
             "cause is uncertainty about demand, and simple ordering rules ignore the things that drive it.")

    # 3 ---------------------------------------------------------------- objectives
    n += 1
    s = content_slide(prs, "Objectives", "Four objectives, each building on the last", n + 1)
    import pandas as pd
    rec = pd.read_csv(ROOT / "outputs" / "tables" / "reorder_recommendations.csv").sort_values("stockout_risk", ascending=False)
    top = rec.iloc[0]
    upl = R["uplift"].set_index("segment_name")["mean"]
    objs = [("1", "Predict demand", "Daily sales per store for the next 7 days, with a P10-P90 range",
             "e.g. store 262, Mon 1 Jun: 24,078", NAVY),
            ("2", "Predict stockout risk", "Probability of running out in the next 3 days",
             f"e.g. store {int(top.Store)}: {top.stockout_risk:.0%} → {top.risk_level}", RED),
            ("3", "Recommend the order", "Order-up-to level from forecast + uncertainty; emergency order if risk is high",
             f"e.g. order up to {top.order_up_to:,.0f}", AMBER),
            ("4", "Explain + what-if", "SHAP reasons for every number; test promotion, holiday, competitor, price",
             f"e.g. promotion: +{upl.min() * 100:.0f}% to +{upl.max() * 100:.0f}%", TEAL)]
    for i, (num, head, body, ex, col) in enumerate(objs):
        x = 0.6 + i * 3.1
        box(s, x, 1.75, 2.9, 4.35, fill=TINT)
        circle(s, x + 0.3, 2.0, 0.7, col, num, size=24)
        text(s, x + 0.3, 2.95, 2.4, 0.8, head, size=20, bold=True, font=HEAD, color=INK)
        text(s, x + 0.3, 3.75, 2.35, 1.4, body, size=15, color=INK)
        box(s, x + 0.2, 5.15, 2.5, 0.75, fill=WHITE)
        text(s, x + 0.35, 5.2, 2.25, 0.65, ex, size=13, color=col, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.6, 6.35, 12, 0.4, "Supporting objective: group the 935 stores into segments with K-Means so one plan does not have to fit all.",
         size=15, color=MUTED, italic=True)
    notes(s, "Objective one is the forecast. Objective two uses it with current stock to give a risk probability. "
             "Objective three turns both into an order quantity. Objective four explains each number and lets the user "
             "play what-if. Clustering is a supporting objective.")

    # 4 ---------------------------------------------------------------- literature
    n += 1
    s = content_slide(prs, "Literature survey", "What 20 books and 15 research papers told us", n + 1)
    for i, (big, lab, col) in enumerate([("20", "reference books\nforecasting, inventory theory,\nML, explainability, causality", NAVY),
                                         ("15", "research papers\nmostly 2019 - 2026,\nIJF, EJOR, NeurIPS, IEEE Access", TEAL)]):
        y = 1.75 + i * 2.45
        box(s, 0.6, y, 3.6, 2.2, fill=col)
        text(s, 0.9, y + 0.2, 3.0, 1.0, big, size=54, bold=True, color=WHITE, font=HEAD)
        text(s, 0.9, y + 1.15, 3.1, 0.95, lab, size=13, color=WHITE)
    rows = [("ML beats statistics on retail data", "M5 competition: LightGBM-based methods won on Walmart data", "Makridakis et al., 2022"),
            ("Promotions and special days need explicit features", "Regression trees with promo flags; calendar special days", "Gür Ali 2009; Huber 2020"),
            ("Stockout prediction is rare - and needs stock + forecast", "1.6 M SKUs: stock level and near-term forecast matter most", "Liu et al., 2025"),
            ("A more accurate forecast is not always a better decision", "Weak link between accuracy and inventory cost on M5", "Theodorou et al., 2025"),
            ("Explanations stop at the forecast", "LightGBM + SHAP/ICE explain demand, not risk or orders", "Mishra et al., 2026")]
    for i, (head, body, src) in enumerate(rows):
        y = 1.75 + i * 0.9
        circle(s, 4.6, y + 0.12, 0.5, AMBER, str(i + 1), size=15)
        text(s, 5.3, y + 0.05, 7.4, 0.35, head, size=17, bold=True, color=INK)
        text(s, 5.3, y + 0.43, 7.4, 0.35, [[(body + "   ", {}), (src, {"italic": True, "color": GREY, "size": 12})]],
             size=14, color=MUTED)
    notes(s, "We read 20 textbooks for foundations and 15 papers for the state of the art. Five lessons shaped the "
             "design: ML wins on retail data, promotions need features, stockout prediction needs stock plus forecast, "
             "accuracy is not the same as a good decision, and explanations usually stop at the forecast.")

    # 5 ---------------------------------------------------------------- gaps
    n += 1
    s = content_slide(prs, "Research gaps", "Six gaps - and how DemandIQ answers each one", n + 1)
    gaps = [("G1", "Forecast, risk and ordering studied separately", "One pipeline: forecast → risk → order"),
            ("G2", "Explanations end at the forecast", "SHAP for the demand AND the risk model"),
            ("G3", "Accuracy is not decision quality", "Judge on stockouts, fill rate, stock and cost"),
            ("G4", "What-if read as cause and effect", "Every what-if labelled as model behaviour"),
            ("G5", "One policy for very different stores", "K-Means segments; uplift per segment"),
            ("G6", "No public data has stock on hand", "Calibrated inventory replay over real sales")]
    for i, (gid, gap, ans) in enumerate(gaps):
        col, row = i % 2, i // 2
        x, y = 0.6 + col * 6.15, 1.75 + row * 1.6
        box(s, x, y, 5.95, 1.4, fill=TINT)
        circle(s, x + 0.25, y + 0.38, 0.64, NAVY, gid, size=14)
        text(s, x + 1.1, y + 0.2, 4.7, 0.5, gap, size=16, bold=True, color=INK)
        text(s, x + 1.1, y + 0.75, 4.7, 0.5, [[("→  ", {"color": TEAL, "bold": True}), (ans, {})]], size=15, color=TEAL)
    notes(s, "Each gap links back to specific papers in the report. G1 and G2 are the main contribution: one "
             "pipeline that is explained end to end. G6 is a practical gap: public data does not contain stock levels.")

    # 6 ---------------------------------------------------------------- data
    n += 1
    s = content_slide(prs, "Dataset", "Real sales from 1,115 Rossmann drugstores, 2013 - 2015", n + 1)
    stats = [(f"{ds['raw_rows'] / 1e6:.2f} M", "daily store records"), ("942", "days of history"),
             (f"{ds['stores_kept']}", "stores with complete history"), ("27", "leak-free features")]
    for i, (big, lab) in enumerate(stats):
        x = 0.6 + i * 3.1
        box(s, x, 1.75, 2.9, 1.65, fill=TINT)
        text(s, x + 0.3, 1.9, 2.4, 0.8, big, size=36, bold=True, color=NAVY, font=HEAD)
        text(s, x + 0.3, 2.75, 2.5, 0.5, lab, size=15, color=MUTED)
    text(s, 0.6, 3.75, 6, 0.4, "Split strictly by date - never shuffled", size=18, bold=True, color=INK, font=HEAD)
    segs = [("Training", "Jan 2013 - Mar 2015", 9.3, NAVY), ("Validation", "Apr - May", 1.4, AMBER), ("Test", "Jun - Jul 2015", 1.45, RED)]
    x = 0.6
    for lab, per, w, col in segs:
        box(s, x, 4.3, w - 0.05, 0.8, fill=col, radius=0.15)
        text(s, x + 0.2, 4.36, w - 0.3, 0.35, lab, size=15, bold=True, color=WHITE)
        text(s, x + 0.2, 4.68, w - 0.3, 0.35, per, size=12, color=WHITE)
        x += w
    text(s, 0.6, 5.45, 12.1, 1.3, [
        [("Why Rossmann? ", {"bold": True, "color": NAVY}),
         ("Real retail data with exactly the drivers we need - promotions, public and school holidays, weekday, season, "
          "competition and store format. Public, free and a well-known benchmark.", {})],
        [("Cleaning: ", {"bold": True, "color": NAVY}),
         ("dropped 180 stores with a 6-month refurbishment gap, re-indexed every store to a full calendar, every sales "
          "feature is at least 7 days old.", {})]], size=15, spacing=8)
    notes(s, "The data is public Rossmann data. We keep 935 stores with complete history. The split is by time: train "
             "to March 2015, validate April-May, test June-July. The test window is used once, at the end.")

    # 7 ---------------------------------------------------------------- EDA
    n += 1
    s = content_slide(prs, "Exploratory analysis", f"Promotions add {ds['promo_lift_chain'] * 100:.0f}% - and timing matters", n + 1)
    picture(s, FIG / "fig02_weekday_promo.png", 0.6, 1.65, w=5.7, h=3.4)
    picture(s, FIG / "fig03_seasonality_calendar.png", 6.5, 1.65, w=6.25, h=3.4)
    for i, (head, body) in enumerate([
            ("Promotions", f"run on {ds['promo_share'] * 100:.0f}% of trading days and lift sales by about {ds['promo_lift_chain'] * 100:.0f}%"),
            ("Weekly rhythm", "Monday is the busiest day; most stores close on Sunday"),
            ("Calendar", "December peak; extra sales in school holidays and before public holidays")]):
        x = 0.6 + i * 4.1
        box(s, x, 5.3, 3.9, 1.45, fill=TINT)
        text(s, x + 0.25, 5.42, 3.4, 0.4, head, size=17, bold=True, color=NAVY)
        text(s, x + 0.25, 5.82, 3.45, 0.85, body, size=14, color=INK)
    notes(s, "Three patterns drive demand: promotions, the weekly rhythm and the calendar. A rule that only looks at "
             "last month's average misses all three.")

    # 8 ---------------------------------------------------------------- architecture
    n += 1
    s = content_slide(prs, "System design", "Eight steps from raw sales to an explained order", n + 1)
    picture(s, FIG / "fig00_architecture.png", 0.6, 1.55, w=12.15, h=5.35)
    notes(s, "Each box is one Python script. Clean and build features, segment stores, forecast demand, replay "
             "inventory and predict risk, decide the order, explain with SHAP, and show it in the dashboard.")

    # 9 ---------------------------------------------------------------- segments
    n += 1
    s = content_slide(prs, "Unsupervised learning", "K-Means finds four kinds of stores", n + 1)
    picture(s, FIG / "fig06_segment_profiles.png", 0.5, 1.6, w=7.6, h=3.6)
    seg = R["seg"].set_index("segment_name")
    personas = [("High-volume flagships", NAVY, f"{seg.loc['High-volume flagships', 'avg_sales']:,.0f}/day, steady, close competitors"),
                ("Promotion-driven", TEAL, f"+{seg.loc['Promotion-driven', 'promo_lift'] * 100:.0f}% on promotion days, most volatile"),
                ("Weekend & seasonal", AMBER, "strong Saturdays and December"),
                ("Steady neighbourhood", RED, "close to average on every trait")]
    for i, (nm, col, d) in enumerate(personas):
        y = 1.65 + i * 0.95
        circle(s, 8.45, y + 0.12, 0.5, col, str(i + 1), size=14)
        text(s, 9.1, y + 0.05, 3.7, 0.35, f"{nm}  ({int(seg.loc[[k for k in seg.index if k.startswith(nm.split()[0])][0], 'stores'])})",
             size=16, bold=True, color=INK)
        text(s, 9.1, y + 0.42, 3.7, 0.4, d, size=13, color=MUTED)
    box(s, 0.6, 5.55, 12.15, 1.15, fill=TINT)
    text(s, 0.9, 5.65, 11.6, 0.95, [[("Why it matters: ", {"bold": True, "color": NAVY}),
                                    ("the same promotion lifts a promotion-driven store far more than a flagship, so one "
                                     "promotion plan and one safety buffer do not fit all stores. The segment is also a feature "
                                     "for both models.", {})]], size=15, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "We cluster stores on six traits measured on the training window. Silhouette scores are low for every k, "
             "so stores form a continuum; we chose k equals 4 because the elbow flattens and the groups are easy to act on.")

    # 10 --------------------------------------------------------------- demand results (native chart)
    n += 1
    s = content_slide(prs, "Demand forecasting", f"Gradient Boosting cuts forecast error to {gb['MAPE (%)']:.1f}%", n + 1)
    order = ["Naive (same day last week)", "Historical average", "Ridge regression", "Random Forest", "Gradient Boosting"]
    labels = ["Naive (last week)", "Historical average", "Ridge regression", "Random Forest", "Gradient Boosting"]
    cd = CategoryChartData()
    cd.categories = labels
    cd.add_series("Test MAPE (%)", [round(dt.loc[m, "MAPE (%)"], 2) for m in order])
    gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.6), Inches(7.6), Inches(5.1), cd)
    ch = gf.chart
    style_chart(ch, 13)
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = "Test MAPE, June-July 2015 (lower is better)"
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = INK
    plot = ch.plots[0]
    plot.gap_width = 55
    plot.has_data_labels = True
    plot.data_labels.number_format = '0.0"%"'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size = Pt(13)
    plot.data_labels.font.bold = True
    cols = [GREY, GREY, TEAL, TEAL, NAVY]
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = cols[i]
    ch.category_axis.reverse_order = True
    ch.category_axis.tick_labels.font.size = Pt(13)
    ch.category_axis.format.line.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE3, 0xE3)
    va.maximum_scale = 36
    va.minimum_scale = 0
    va.major_unit = 10
    from pptx.enum.chart import XL_TICK_LABEL_POSITION
    va.tick_label_position = XL_TICK_LABEL_POSITION.NONE
    va.format.line.fill.background()
    for i, (big, lab) in enumerate([(f"{gb['R2']:.3f}", "R² on the test window"),
                                    (f"{gb['Bias (%)']:+.1f}%", "bias - almost no systematic over- or under-forecast"),
                                    (f"{R['q']['p10_p90_coverage_test'] * 100:.0f}%", "of days inside the P10-P90 range (target 80%)")]):
        y = 1.65 + i * 1.7
        box(s, 8.6, y, 4.15, 1.5, fill=TINT)
        text(s, 8.9, y + 0.12, 3.6, 0.7, big, size=32, bold=True, color=NAVY, font=HEAD)
        text(s, 8.9, y + 0.8, 3.65, 0.6, lab, size=14, color=MUTED)
    notes(s, f"The naive forecast is poor because promotions alternate weekly. The historical average is a strong "
             f"baseline at {ha['MAPE (%)']:.1f}%. Ridge is worse because it cannot model interactions. Gradient Boosting was "
             f"best on validation and reaches {gb['MAPE (%)']:.2f}% on the test window, R-squared {gb['R2']:.3f}.")

    # 11 --------------------------------------------------------------- forecast examples
    n += 1
    s = content_slide(prs, "Demand forecasting", "The forecast follows promotions and Monday peaks", n + 1)
    picture(s, FIG / "fig10_forecast_examples.png", 0.6, 1.5, w=8.6, h=5.4)
    for i, (head, body) in enumerate([
            ("Leak-free by design", "Every sales feature is at least 7 days old, so one model forecasts the whole coming week."),
            ("Uncertainty, not just a number", "Quantile models give P10 and P90; the band feeds the safety stock."),
            ("Honest tuning", "Grid search on Apr-May; the test window was opened only once.")]):
        y = 1.65 + i * 1.75
        box(s, 9.5, y, 3.25, 1.55, fill=TINT)
        text(s, 9.75, y + 0.12, 2.85, 0.4, head, size=16, bold=True, color=NAVY)
        text(s, 9.75, y + 0.52, 2.85, 1.0, body, size=13, color=INK)
    notes(s, "Black dots are actual sales, the orange line is our forecast and the band is the P10 to P90 range. Red "
             "squares mark promotion days, and the forecast jumps with them.")

    # 12 --------------------------------------------------------------- inventory replay
    n += 1
    s = content_slide(prs, "Creating stock data", "Rossmann has no stock data - so we replayed it", n + 1)
    picture(s, FIG / "fig12_inventory_replay.png", 0.6, 1.55, w=7.9, h=3.4)
    steps = [("Morning", "deliveries arrive - 20% of them late", NAVY),
             ("Day", "customers buy; if the shelf is empty the sale is lost", RED),
             ("Evening", "Mon & Thu: order up to a target level", AMBER)]
    for i, (h_, b_, col) in enumerate(steps):
        y = 1.65 + i * 1.1
        circle(s, 8.85, y + 0.08, 0.6, col, str(i + 1), size=16)
        text(s, 9.6, y + 0.05, 3.2, 0.35, h_, size=17, bold=True, color=INK)
        text(s, 9.6, y + 0.42, 3.2, 0.6, b_, size=13, color=MUTED)
    box(s, 0.6, 5.25, 12.15, 1.45, fill=NAVY)
    text(s, 0.95, 5.35, 2.4, 1.2, f"{inv['stockout_day_rate_%']:.1f}%", size=40, bold=True, color=WHITE, font=HEAD,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, 3.35, 5.35, 9.1, 1.2, [[("stockout days under the manager's rule (last 4 weeks' average + 20%). ", {"bold": True}),
                                   ("We calibrated the buffer so the replay matches the 8.3% out-of-stock rate measured in real "
                                    f"stores. Promotion days: {inv['stockout_rate_promo_days_%']:.0f}% stockouts, other days "
                                    f"{inv['stockout_rate_non_promo_%']:.0f}%.", {})]],
         size=15, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "No public dataset has stock levels. We replay a realistic ordering process over the real sales, with "
             "late deliveries, and calibrate it to the 8.3% stockout rate found in real stores. This gives stock levels "
             "and stockout labels.")

    # 13 --------------------------------------------------------------- risk model
    n += 1
    s = content_slide(prs, "Stockout risk (classification)", "Spotting a stockout three days before it happens", n + 1)
    picture(s, FIG / "fig13_roc_pr_curves.png", 0.6, 1.55, w=7.6, h=3.25)
    picture(s, FIG / "fig15_risk_bands.png", 0.75, 4.85, w=5.0, h=2.1)
    for i, (big, lab, col) in enumerate([
            (f"{gbs['ROC-AUC']:.3f}", f"ROC-AUC (rule of thumb: {st.loc['Manager rule of thumb', 'ROC-AUC']:.2f})", NAVY),
            (f"{gbs['PR-AUC']:.2f}", f"PR-AUC - over 3x the {R['cm']['test_positive_rate'] * 100:.0f}% base rate", TEAL),
            (f"{gbs['Recall'] * 100:.0f}%", "of stockouts caught at the tuned threshold", RED)]):
        y = 1.65 + i * 1.7
        box(s, 8.6, y, 4.15, 1.5, fill=TINT)
        text(s, 8.9, y + 0.12, 3.6, 0.7, big, size=32, bold=True, color=col, font=HEAD)
        text(s, 8.9, y + 0.8, 3.65, 0.6, lab, size=14, color=MUTED)
    b = R["bands"]
    text(s, 5.95, 5.15, 2.4, 1.6, [[("HIGH", {"bold": True, "color": RED}), (f" stores really ran out {b.rate.iloc[2]:.0%} of the time; ", {}),
                                    ("LOW", {"bold": True, "color": GREEN}), (f" only {b.rate.iloc[0]:.0%}.", {})]], size=13, color=INK)
    notes(s, "The classifier sees stock, deliveries due, the forecast and promotions ahead. The manager's rule of thumb "
             "is barely better than chance because it ignores promotions and late deliveries. Our model reaches ROC-AUC "
             f"{gbs['ROC-AUC']:.3f}. Risk is shown as LOW, MEDIUM or HIGH.")

    # 14 --------------------------------------------------------------- policy comparison (native chart)
    n += 1
    s = content_slide(prs, "From prediction to decision", f"{abs(C['stockout_change_%']):.0f}% fewer stockout days, "
                      f"{abs(C['cost_change_%']):.0f}% lower cost", n + 1)
    cd = CategoryChartData()
    cd.categories = ["A. Manager rule", "B. DemandIQ forecast", "C. + risk alerts"]
    cd.add_series("Stockout days (%)", [round(x, 2) for x in pol["stockout_day_rate_%"]])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.6), Inches(1.6), Inches(5.9), Inches(4.0), cd)
    ch = gf.chart
    style_chart(ch, 13)
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = "Stockout days, test window (%)"
    r0 = ch.chart_title.text_frame.paragraphs[0].runs[0]
    r0.font.size, r0.font.bold, r0.font.color.rgb = Pt(14), True, INK
    plot = ch.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    plot.data_labels.number_format = '0.00"%"'
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size = Pt(14)
    plot.data_labels.font.bold = True
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = [GREY, TEAL, NAVY][i]
    ch.value_axis.has_major_gridlines = True
    ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xE3, 0xE3, 0xE3)
    ch.value_axis.format.line.fill.background()
    ch.value_axis.maximum_scale = 9
    ch.category_axis.tick_labels.font.size = Pt(12)
    picture(s, FIG / "fig17_cost_breakdown.png", 6.75, 1.6, w=6.0, h=3.0)
    rows = [("Policy", "Fill rate", "Stock (days)", "Cost vs A")] + \
           [(p_.split(". ")[0], f"{r['fill_rate_%']:.2f}%", f"{r['avg_stock_days']:.2f}", f"{r['cost_change_%']:+.0f}%")
            for p_, (_, r) in zip(pol.policy, pol.iterrows())]
    tbl = s.shapes.add_table(4, 4, Inches(6.9), Inches(4.75), Inches(5.85), Inches(1.6)).table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            cell = tbl.cell(i, j)
            cell.text = v
            p = cell.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            p.runs[0].font.size = Pt(13)
            p.runs[0].font.name = BODY
            p.runs[0].font.bold = i == 0 or i == 3
            p.runs[0].font.color.rgb = WHITE if i == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if i == 0 else (TINT if i % 2 == 0 else WHITE)
    box(s, 0.6, 5.85, 5.9, 0.85, fill=TINT)
    text(s, 0.85, 5.9, 5.5, 0.75, "A better forecast removes most stockouts; risk-triggered emergency orders remove much of the rest.",
         size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, f"All three policies face the same real demand and the same late deliveries. The manager rule gives "
             f"{A['stockout_day_rate_%']:.2f}% stockout days. Using our forecast gives {B['stockout_day_rate_%']:.2f}%. Adding "
             f"risk-triggered emergency orders gives {C['stockout_day_rate_%']:.2f}%, and total cost falls by "
             f"{abs(C['cost_change_%']):.0f}%. Lost margin is the biggest cost for the manager.")

    # 15 --------------------------------------------------------------- frontier
    n += 1
    s = content_slide(prs, "Fair comparison", "Same stock, fewer stockouts - it is not just \"buy more\"", n + 1)
    picture(s, FIG / "fig24_tradeoff_frontier.png", 0.6, 1.55, w=7.3, h=5.2)
    for i, (head, body) in enumerate([
            ("The question", "Could the manager do just as well by keeping a bigger buffer?"),
            ("The test", "Replay both policies at many buffer and service-level settings."),
            ("The answer", "At every level of stock, DemandIQ has fewer stockout days - it puts the stock on the right "
                           "days, before promotions and holidays.")]):
        y = 1.65 + i * 1.7
        box(s, 8.3, y, 4.45, 1.5, fill=TINT if i < 2 else NAVY)
        text(s, 8.55, y + 0.12, 4.0, 0.4, head, size=16, bold=True, color=NAVY if i < 2 else WHITE)
        text(s, 8.55, y + 0.5, 4.0, 1.0, body, size=13, color=INK if i < 2 else WHITE)
    notes(s, "This is our answer to the obvious objection. Each dot is one setting. The DemandIQ curve is below the "
             "manager's curve everywhere: for the same stock, fewer stockouts.")

    # 16 --------------------------------------------------------------- SHAP
    n += 1
    s = content_slide(prs, "Explainability (SHAP)", "Every number comes with its reasons", n + 1)
    picture(s, FIG / "fig19_shap_demand_global.png", 0.6, 1.5, w=5.0, h=4.0)
    picture(s, FIG / "fig21_shap_demand_local.png", 5.85, 1.4, w=6.9, h=4.15)
    text(s, 0.6, 5.55, 5.0, 0.35, "Global: what drives demand across all stores", size=14, bold=True, color=NAVY)
    text(s, 5.9, 5.55, 6.85, 0.35, "Local: why store 262 got 24,078 on Mon 1 Jun (actual 23,204)", size=14, bold=True, color=NAVY)
    box(s, 0.6, 6.0, 12.15, 0.75, fill=TINT)
    up = R["uplift"].set_index("segment_name")["mean"]
    text(s, 0.85, 6.03, 11.7, 0.7, [[("What-if with the model: ", {"bold": True, "color": NAVY}),
                                    (f"switching a promotion on lifts the forecast by {up.min() * 100:.0f}% to {up.max() * 100:.0f}% "
                                     "depending on the store segment. Shown as model behaviour, not proof of cause and effect.", {})]],
         size=14, anchor=MSO_ANCHOR.MIDDLE)
    notes(s, "SHAP splits each prediction into the push from each input. On the left, the global view. On the right, "
             "one forecast: start from the average day, and each bar pushes it up or down. SHAP values here are exact "
             "and add up to the prediction.")

    # 17 --------------------------------------------------------------- dashboard
    n += 1
    s = content_slide(prs, "Dashboard", "A six-page Streamlit app that managers can use", n + 1)
    shots = [p for p in (SHOTS / "whatif.png", SHOTS / "risk.png") if p.exists()]
    for i, p in enumerate(shots):
        im = Image.open(p)
        crop = im.crop((int(im.width * 0.205), int(im.height * 0.045), im.width, int(im.height * (0.86 if i == 0 else 0.6))))
        cp = ROOT / "_qa" / f"crop_{p.stem}.png"
        crop.save(cp)
        box(s, 0.6 + i * 6.15, 1.55, 5.95, 4.4, fill=TINT)
        picture(s, cp, 0.7 + i * 6.15, 1.65, w=5.75, h=4.2)
    text(s, 0.6, 6.1, 12.15, 0.6, "Overview  •  Store forecast  •  Risk & reorder list  •  Why this forecast?  "
         "•  What-if simulator  •  Store segments", size=15, color=MUTED, align=PP_ALIGN.CENTER)
    notes(s, "Live demo: pick a store, switch on a promotion in the what-if page, and see the forecast and order change. "
             "The risk page lists stores by risk with a recommended order.")

    # 18 --------------------------------------------------------------- limitations & future
    n += 1
    s = content_slide(prs, "Limitations and future scope", "What we would do next", n + 1)
    lim = ["Stock is simulated, not observed", "Store level, not product (SKU) level",
           "No price data - price lever uses a stated elasticity",
           f"P10-P90 band covers {R['q']['p10_p90_coverage_test'] * 100:.0f}%, not 80%"]
    fut = ["Validate on data with real stockout labels (FreshRetailNet-50K)", "Indian retail with a festival calendar",
           "Causal price effect with double/debiased ML", "Weekly retraining and an evening HIGH-risk alert"]
    for i, (head, items, col) in enumerate([("Limitations", lim, RED), ("Future scope", fut, TEAL)]):
        x = 0.6 + i * 6.15
        box(s, x, 1.65, 5.95, 5.0, fill=TINT)
        text(s, x + 0.35, 1.85, 5.3, 0.5, head, size=21, bold=True, color=col, font=HEAD)
        for k, it in enumerate(items):
            y = 2.6 + k * 0.98
            circle(s, x + 0.35, y + 0.05, 0.42, col, str(k + 1), size=13)
            text(s, x + 0.95, y, 4.75, 0.85, it, size=16, color=INK, anchor=MSO_ANCHOR.TOP)
    notes(s, "We state limitations openly. The biggest is that stock levels are simulated; the next step is a dataset "
             "with real stockout labels and product-level demand.")

    # 19 --------------------------------------------------------------- conclusion
    n += 1
    s = prs.slides.add_slide(prs.slide_layouts[6])
    background(s, NAVY)
    text(s, 0.6, 0.6, 12, 0.4, "CONCLUSION", size=12, bold=True, color=RGBColor(0xF2, 0xC1, 0x6B))
    text(s, 0.6, 0.95, 12, 0.9, "A forecast creates value only when it becomes an explained decision", size=30,
         bold=True, color=WHITE, font=HEAD)
    for i, (big, lab) in enumerate([(f"{gb['MAPE (%)']:.1f}%", "demand forecast error (MAPE)"),
                                    (f"{gbs['ROC-AUC']:.2f}", "ROC-AUC for 3-day stockout risk"),
                                    (f"{abs(C['stockout_change_%']):.0f}%", "fewer stockout days than the manager rule"),
                                    (f"{abs(C['cost_change_%']):.0f}%", "lower total inventory cost")]):
        x = 0.6 + i * 3.1
        box(s, x, 2.3, 2.9, 1.9, fill=NAVY_DK)
        text(s, x + 0.3, 2.45, 2.4, 0.9, big, size=40, bold=True, color=WHITE, font=HEAD)
        text(s, x + 0.3, 3.35, 2.45, 0.7, lab, size=14, color=RGBColor(0xC9, 0xD6, 0xE3))
    text(s, 0.6, 4.6, 12, 0.8, "Classical ML only  •  Random Forest, Gradient Boosting, Logistic Regression, K-Means  "
         "•  SHAP  •  Streamlit", size=16, color=RGBColor(0xC9, 0xD6, 0xE3))
    text(s, 0.6, 5.6, 12, 0.9, "Thank you  -  questions?", size=36, bold=True, color=RGBColor(0xF2, 0xC1, 0x6B), font=HEAD)
    text(s, 0.6, 6.6, 12.1, 0.4, "   |   ".join(f"{nm} ({r})" for nm, r in TEAM), size=13, color=RGBColor(0xC9, 0xD6, 0xE3))
    notes(s, "To summarise: an accurate forecast, a calibrated risk score, a transparent order rule and SHAP "
             "explanations together cut stockout days sharply at lower cost. Thank you, we are happy to take questions.")

    out = HERE / "DemandIQ_Presentation.pptx"
    prs.core_properties.title = "DemandIQ"
    prs.core_properties.author = ", ".join(nm for nm, _ in TEAM)
    prs.save(out)
    print("saved", out, "| slides:", len(prs.slides))


if __name__ == "__main__":
    build()
