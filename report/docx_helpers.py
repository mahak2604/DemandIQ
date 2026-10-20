"""Formatting helpers for the DemandIQ report (python-docx)."""
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

FONT = "Times New Roman"
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
HEADER_FILL = "1F3A5F"
ZEBRA_FILL = "EEF2F6"


def new_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)          # A4
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)

    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = st.paragraph_format
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.0

    for name, size, before, after in (("Heading 1", 16, 12, 8), ("Heading 2", 13, 10, 5),
                                      ("Heading 3", 12, 8, 4)):
        h = doc.styles[name]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.italic = False
        h.font.color.rgb = NAVY if name != "Heading 3" else RGBColor(0x22, 0x22, 0x22)
        rpr = h.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(a), FONT)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rfonts.get(qn(a)) is not None:
                del rfonts.attrib[qn(a)]
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True

    cap = doc.styles["Caption"]
    cap.font.name = FONT
    cap.font.size = Pt(10)
    cap.font.bold = False
    cap.font.italic = False
    cap.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    from docx.enum.style import WD_STYLE_TYPE
    if "Table Caption" not in [x.name for x in doc.styles]:
        tc = doc.styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
        tc.base_style = cap
        tc.font.name = FONT
        tc.font.size = Pt(10)
        tc.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    for lst in ("List Bullet", "List Number"):
        s = doc.styles[lst]
        s.font.name = FONT
        s.font.size = Pt(12)
        s.paragraph_format.space_after = Pt(3)
    return doc


# ---------------------------------------------------------------- text
def para(doc, text="", bold=False, italic=False, size=None, align="justify", after=6, before=0,
         color=None, keep_next=False, indent=None):
    p = doc.add_paragraph()
    add_runs(p, text, bold=bold, italic=italic, size=size, color=color)
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if keep_next:
        p.paragraph_format.keep_with_next = True
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
    return p


def add_runs(p, text, bold=False, italic=False, size=None, color=None):
    """Supports **bold** and _italic_ spans inside text."""
    import re
    parts = re.split(r"(\*\*[^*]+\*\*|(?<![A-Za-z0-9])_[^_]+?_(?![A-Za-z0-9]))", text)
    for part in parts:
        if not part:
            continue
        b, i, t = bold, italic, part
        if part.startswith("**") and part.endswith("**"):
            b, t = True, part[2:-2]
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            i, t = True, part[1:-1]
        r = p.add_run(t)
        r.bold, r.italic = b, i
        r.font.name = FONT
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def bullets(doc, items, style="List Bullet"):
    for it in items:
        p = doc.add_paragraph(style=style)
        add_runs(p, it)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(3)


def equation(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Cambria Math"
    r.font.size = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    return p


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- figures & tables
class Counter:
    fig = 0
    tab = 0


def figure(doc, path, caption, width=6.2):
    Counter.fig += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    c = doc.add_paragraph(style="Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Fig. {Counter.fig}. ")
    r.bold = True
    add_runs(c, caption)
    c.paragraph_format.space_after = Pt(10)
    return Counter.fig


def next_fig():
    return Counter.fig + 1


def next_tab():
    return Counter.tab + 1


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_margins(table, top=50, bottom=50, left=80, right=80):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    e.set(qn("w:val"), "true")
    trPr.append(e)


def table(doc, caption, header, rows, widths, font_size=10, align=None, bold_rows=(), zebra=True):
    """Caption above (IEEE style), navy header row, light zebra rows."""
    Counter.tab += 1
    c = doc.add_paragraph(style="Table Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Table {Counter.tab}. ")
    r.bold = True
    add_runs(c, caption)
    c.paragraph_format.keep_with_next = True
    c.paragraph_format.space_before = Pt(6)
    c.paragraph_format.space_after = Pt(3)

    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_cell_margins(t)
    align = align or ["left"] + ["center"] * (len(header) - 1)
    amap = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.LEFT}
    for ri, vals in enumerate([header] + rows):
        row = t.rows[ri]
        no_split(row)
        if ri == 0:
            repeat_header(row)
        for ci, v in enumerate(vals):
            cell = row.cells[ci]
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if len(rows) <= 10 and ri < len(rows):
                p.paragraph_format.keep_with_next = True      # keep short tables on one page
            p.alignment = amap[align[ci]] if ri > 0 else WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, str(v), bold=(ri == 0 or (ri - 1) in bold_rows), size=font_size,
                     color=RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else None)
            if ri == 0:
                shade(cell, HEADER_FILL)
            elif zebra and ri % 2 == 0:
                shade(cell, ZEBRA_FILL)
    # grid widths
    grid = t._tbl.tblGrid
    for gi, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[gi] * 1440)))
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)
    return t


# ---------------------------------------------------------------- fields
def field(paragraph, instr):
    r = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r._r.append(b)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = paragraph.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r3._r.append(s)
    r4 = paragraph.add_run("(right-click and choose Update Field)")
    r5 = paragraph.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r5._r.append(e)
    return paragraph


def page_number_footer(section, left_text):
    section.footer.is_linked_to_previous = False
    f = section.footer.paragraphs[0]
    f.text = ""
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = f.add_run(left_text + "   |   Page ")
    r.font.size = Pt(9)
    r.font.name = FONT
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    run = f.add_run()
    for kind, txt in (("begin", None), ("instr", "PAGE"), ("separate", None), ("text", "1"), ("end", None)):
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = txt
            run._r.append(it)
        elif kind == "text":
            t = OxmlElement("w:t")
            t.text = txt
            run._r.append(t)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            run._r.append(fc)
    run.font.size = Pt(9)
    run.font.name = FONT


def header_text(section, text):
    section.header.is_linked_to_previous = False
    h = section.header.paragraphs[0]
    h.text = ""
    h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = h.add_run(text)
    r.font.size = Pt(9)
    r.italic = True
    r.font.name = FONT
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    # thin rule under the header
    pPr = h._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single")
    bot.set(qn("w:sz"), "4")
    bot.set(qn("w:space"), "1")
    bot.set(qn("w:color"), "999999")
    bdr.append(bot)
    pPr.append(bdr)


def page_number_format(section, fmt="decimal", start=1):
    sectPr = section._sectPr
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), str(start))
    sectPr.append(pg)


def page_border(section):
    sectPr = section._sectPr
    b = OxmlElement("w:pgBorders")
    b.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:val"), "double")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "24")
        e.set(qn("w:color"), "1F3A5F")
        b.append(e)
    sectPr.append(b)
